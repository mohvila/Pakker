#!/usr/bin/env python3
"""Ugentlig prissammenligning for Luxury-Outdoor.dk.

1. Henter alle produkter fra webshoppen (Shopify, WooCommerce eller sitemap + JSON-LD).
2. Slår hvert produkt op på nettet (Google Shopping via SerpApi og/eller PriceRunner).
3. Laver en rapport sorteret efter leverandør, med Lechuza først.
4. Gemmer rapporten som Markdown, Excel, CSV og HTML i OUT_DIR (seneste.* + arkiv/<dato>.*).
   Er SMTP_HOST sat, sendes den desuden på mail.

Indstillinger læses fra miljøvariabler, se README.md i samme mappe.
"""

import csv
import datetime as dt
import difflib
import html
import io
import json
import os
import re
import smtplib
import ssl
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from email.message import EmailMessage

SHOP_URL = os.environ.get("SHOP_URL", "https://luxury-outdoor.dk").rstrip("/")
SHOP_HOST = urllib.parse.urlparse(SHOP_URL).hostname or ""
FIRST_SUPPLIERS = [s.strip() for s in os.environ.get("FIRST_SUPPLIERS", "Lechuza").split(",") if s.strip()]
MAIL_TO = os.environ.get("MAIL_TO", "info@luxury-outdoor.dk")
MAX_PRODUCTS = int(os.environ.get("MAX_PRODUCTS", "0") or 0)
MIN_MATCH = float(os.environ.get("MIN_MATCH", "0.7"))
DELAY = float(os.environ.get("REQUEST_DELAY", "1.0"))
OUT_DIR = os.environ.get("OUT_DIR", os.path.join(os.path.dirname(os.path.abspath(__file__)), "rapporter"))
UA = "Mozilla/5.0 (compatible; LuxuryOutdoorPrisTjek/1.0; +" + SHOP_URL + ")"

# Leverandører der genkendes i produktnavnet, hvis shoppen ikke selv angiver en.
KNOWN_BRANDS = [
    "Lechuza", "Weber", "Kettler", "Cane-line", "Cane line", "Brafab", "Hartman", "Glatz",
    "Tuuci", "Sika-Design", "Sika Design", "Fatboy", "Kartell", "Vondom", "Serralunga",
    "Elho", "Capi", "Pottery Pots", "Umbrosa", "Outdoor Chef", "Big Green Egg", "Ooni",
    "Morsø", "Eva Solo", "Traeger", "Napoleon", "Gardenas", "Unopiu", "Emu", "Fermob",
    "Houe", "Skagerak", "Cinas", "Todus", "Dedon", "Kristalia", "Diphano", "Borek",
]


# ---------------------------------------------------------------- HTTP-hjælpere

def get(url, params=None, accept="application/json", retries=3):
    if params:
        url += ("&" if "?" in url else "?") + urllib.parse.urlencode(params)
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": accept,
                                               "Accept-Language": "da-DK,da;q=0.9"})
    for attempt in range(retries):
        try:
            with urllib.request.urlopen(req, timeout=30) as r:
                return r.read().decode(r.headers.get_content_charset() or "utf-8", "replace")
        except urllib.error.HTTPError as e:
            if e.code in (404, 400, 401, 403):
                raise
            err = e
        except (urllib.error.URLError, TimeoutError) as e:
            err = e
        time.sleep(2 ** (attempt + 1))
    raise err


def get_json(url, params=None):
    return json.loads(get(url, params))


def to_float(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-]", "", str(v))
    if not s:
        return None
    # Dansk format "1.299,95" og engelsk "1299.95"
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def strip_tags(s):
    return html.unescape(re.sub(r"<[^>]+>", "", s or "")).strip()


# ---------------------------------------------------------------- Produkter fra shoppen

def product(name, supplier, price, url, sku="", gtin=""):
    return {"name": strip_tags(name), "supplier": (supplier or "").strip(), "price": price,
            "url": url, "sku": sku or "", "gtin": gtin or ""}


def fetch_shopify():
    items, page = [], 1
    while True:
        data = get_json(f"{SHOP_URL}/products.json", {"limit": 250, "page": page})
        batch = data.get("products", [])
        if not batch:
            break
        for p in batch:
            variants = p.get("variants") or [{}]
            v = min(variants, key=lambda x: to_float(x.get("price")) or 1e12)
            items.append(product(p.get("title"), p.get("vendor"), to_float(v.get("price")),
                                 f"{SHOP_URL}/products/{p.get('handle')}", v.get("sku"), v.get("barcode")))
        page += 1
        time.sleep(DELAY)
    return items


def fetch_woocommerce():
    items, page = [], 1
    while True:
        batch = get_json(f"{SHOP_URL}/wp-json/wc/store/v1/products", {"per_page": 100, "page": page})
        if not batch:
            break
        for p in batch:
            prices = p.get("prices") or {}
            minor = int(prices.get("currency_minor_unit", 2) or 2)
            raw = to_float(prices.get("price"))
            price = raw / (10 ** minor) if raw is not None else None
            brand = ""
            if p.get("brands"):
                brand = p["brands"][0].get("name", "")
            if not brand:
                for a in p.get("attributes") or []:
                    if re.search(r"brand|m[æa]rke|leverand|producent", a.get("name", ""), re.I):
                        terms = a.get("terms") or []
                        brand = terms[0].get("name", "") if terms else ""
                        break
            items.append(product(p.get("name"), brand, price, p.get("permalink"), p.get("sku")))
        page += 1
        time.sleep(DELAY)
    return items


def jsonld_products(page_html):
    for block in re.findall(r'<script[^>]+application/ld\+json[^>]*>(.*?)</script>', page_html, re.S | re.I):
        try:
            data = json.loads(block.strip())
        except ValueError:
            continue
        stack = [data]
        while stack:
            d = stack.pop()
            if isinstance(d, list):
                stack.extend(d)
            elif isinstance(d, dict):
                t = d.get("@type")
                if t == "Product" or (isinstance(t, list) and "Product" in t):
                    yield d
                stack.extend(v for k, v in d.items() if k == "@graph")


def fetch_sitemap():
    urls, seen, todo = [], set(), [f"{SHOP_URL}/sitemap.xml", f"{SHOP_URL}/sitemap_index.xml"]
    while todo:
        sm = todo.pop()
        if sm in seen:
            continue
        seen.add(sm)
        try:
            xml = get(sm, accept="application/xml")
        except Exception:
            continue
        for loc in re.findall(r"<loc>\s*(.*?)\s*</loc>", xml):
            loc = html.unescape(loc)
            if loc.endswith(".xml") or "sitemap" in loc.rsplit("/", 1)[-1]:
                if not re.search(r"(post|page|category|tag|author|blog)", loc, re.I):
                    todo.append(loc)
            else:
                urls.append(loc)
    items = []
    for u in dict.fromkeys(urls):
        if MAX_PRODUCTS and len(items) >= MAX_PRODUCTS:
            break
        try:
            page_html = get(u, accept="text/html")
        except Exception:
            continue
        for d in jsonld_products(page_html):
            offers = d.get("offers") or {}
            if isinstance(offers, list):
                offers = offers[0] if offers else {}
            price = to_float(offers.get("price") or offers.get("lowPrice"))
            brand = d.get("brand") or ""
            if isinstance(brand, dict):
                brand = brand.get("name", "")
            elif isinstance(brand, list):
                brand = brand[0].get("name", "") if brand and isinstance(brand[0], dict) else ""
            gtin = d.get("gtin13") or d.get("gtin") or d.get("gtin12") or d.get("gtin8") or ""
            items.append(product(d.get("name"), brand, price, u, d.get("sku"), gtin))
            break
        time.sleep(DELAY)
    return items


def fetch_products():
    errors = []
    for name, fn in (("Shopify", fetch_shopify), ("WooCommerce", fetch_woocommerce), ("sitemap", fetch_sitemap)):
        try:
            items = fn()
        except Exception as e:
            errors.append(f"{name}: {e}")
            continue
        items = [p for p in items if p["name"]]
        if items:
            print(f"Hentede {len(items)} produkter via {name}", file=sys.stderr)
            return items
    raise SystemExit("Kunne ikke hente produkter fra shoppen: " + "; ".join(errors))


def guess_supplier(p):
    if p["supplier"]:
        return p["supplier"]
    low = p["name"].lower()
    for b in KNOWN_BRANDS:
        if re.search(r"\b" + re.escape(b.lower()) + r"\b", low):
            return b.replace("Cane line", "Cane-line").replace("Sika Design", "Sika-Design")
    return "Ukendt leverandør"


# ---------------------------------------------------------------- Priser på nettet

def norm(s):
    return re.sub(r"[^a-z0-9æøå ]+", " ", (s or "").lower()).split()


def similarity(a, b):
    ta, tb = norm(a), norm(b)
    if not ta or not tb:
        return 0.0
    overlap = len(set(ta) & set(tb)) / len(set(ta))
    ratio = difflib.SequenceMatcher(None, " ".join(ta), " ".join(tb)).ratio()
    # Tal (størrelser, modelnumre) skal passe, ellers er det en anden variant.
    na, nb = set(re.findall(r"\d+", a)), set(re.findall(r"\d+", b))
    if na and not na <= nb:
        return 0.0
    return max(overlap, ratio)


COLORS = {"sort", "hvid", "grå", "grøn", "blå", "rød", "brun", "beige", "antracit", "taupe", "creme", "natur",
          "gul", "orange", "lilla", "pink", "sølv", "guld", "black", "white", "grey", "gray", "green", "blue",
          "red", "brown", "sand", "anthracite", "cream", "silver", "gold"}


def is_identical(p, o):
    """Kun identiske produkter: samme mærke, samme tal (størrelse/model), samme farve og ens navn."""
    title = o["title"]
    tt = set(norm(title))
    brand = norm(p["supplier"])
    if p["supplier"] != "Ukendt leverandør" and brand and not set(brand) <= tt:
        return False
    ours = set(norm(p["name"])) & COLORS
    theirs = tt & COLORS
    if ours and theirs and not ours & theirs:
        return False
    # Søgning på EAN giver sikrere match, så der kræves mindre lighed i navnet.
    return similarity(p["name"], title) >= (0.4 if o.get("exact") else MIN_MATCH)


def is_own_shop(seller, link):
    s = f"{seller} {link}".lower()
    return SHOP_HOST.replace("www.", "") in s or "luxury outdoor" in s or "luxury-outdoor" in s


def search_serpapi(p):
    key = os.environ.get("SERPAPI_KEY")
    if not key:
        return None
    q = p["gtin"] or p["name"]
    data = get_json("https://serpapi.com/search.json", {
        "engine": "google_shopping", "q": q, "gl": "dk", "hl": "da",
        "location": "Denmark", "api_key": key})
    offers = []
    for r in data.get("shopping_results", []) + data.get("inline_shopping_results", []):
        price = r.get("extracted_price") or to_float(r.get("price"))
        link = r.get("product_link") or r.get("link") or ""
        if not price or is_own_shop(r.get("source", ""), link):
            continue
        offers.append({"shop": r.get("source", ""), "price": float(price), "url": link,
                       "title": r.get("title", ""), "exact": bool(p["gtin"])})
    return offers


def walk(obj):
    if isinstance(obj, dict):
        yield obj
        for v in obj.values():
            yield from walk(v)
    elif isinstance(obj, list):
        for v in obj:
            yield from walk(v)


def search_pricerunner(p):
    if os.environ.get("PRICERUNNER", "1") == "0":
        return None
    # PriceRunners offentlige søge-API (samme som hjemmesiden bruger). Uofficiel og kan ændre sig.
    data = get_json("https://www.pricerunner.dk/public/search/v3/DK", {"q": p["name"]})
    offers = []
    for d in walk(data):
        name = d.get("name")
        lp = d.get("lowestPrice") or d.get("price")
        if not name or not isinstance(lp, dict):
            continue
        price = to_float(lp.get("amount"))
        if not price:
            continue
        url = d.get("url") or ""
        if url.startswith("/"):
            url = "https://www.pricerunner.dk" + url
        offers.append({"shop": "PriceRunner (laveste)", "price": price, "url": url, "title": name})
    return offers


def lookup(p):
    offers, errors = [], []
    for fn in (search_serpapi, search_pricerunner):
        try:
            res = fn(p)
        except Exception as e:
            errors.append(f"{fn.__name__.replace('search_', '')}: {e}")
            continue
        if res:
            offers += res
        time.sleep(DELAY)
    matched = [o for o in offers if is_identical(p, o)]
    return matched, errors


# ---------------------------------------------------------------- Rapport

def supplier_key(s):
    low = s.lower()
    for i, first in enumerate(FIRST_SUPPLIERS):
        if low == first.lower():
            return (0, i, low)
    return (2 if s == "Ukendt leverandør" else 1, 0, low)


def build_rows(products):
    rows = []
    for i, p in enumerate(products, 1):
        print(f"[{i}/{len(products)}] {p['supplier']} – {p['name']}", file=sys.stderr)
        offers, errors = lookup(p)
        best = min(offers, key=lambda o: o["price"]) if offers else None
        row = {**p, "best_price": None, "best_shop": "", "best_url": "", "offers": len(offers),
               "diff": None, "diff_pct": None, "status": "Ikke fundet online"}
        if errors and not offers:
            row["status"] = "Fejl ved opslag"
        if best:
            row.update(best_price=best["price"], best_shop=best["shop"], best_url=best["url"])
            if p["price"]:
                row["diff"] = p["price"] - best["price"]
                row["diff_pct"] = row["diff"] / best["price"] * 100
                if row["diff"] > 0.5:
                    row["status"] = "Vi er dyrere"
                elif row["diff"] < -0.5:
                    row["status"] = "Vi er billigst"
                else:
                    row["status"] = "Samme pris"
            else:
                row["status"] = "Mangler vores pris"
        rows.append(row)
    rows.sort(key=lambda r: (supplier_key(r["supplier"]),
                             -(r["diff_pct"] if r["diff_pct"] is not None else -1e9), r["name"].lower()))
    return rows


COLUMNS = [("supplier", "Leverandør"), ("name", "Produkt"), ("sku", "Varenr."), ("gtin", "EAN"),
           ("price", "Vores pris"), ("best_price", "Laveste pris online"), ("best_shop", "Butik"),
           ("diff", "Forskel kr."), ("diff_pct", "Forskel %"), ("offers", "Antal tilbud"),
           ("status", "Status"), ("url", "Vores link"), ("best_url", "Link til tilbud")]


def kr(v):
    return "" if v is None else f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def write_csv(rows):
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow([h for _, h in COLUMNS])
    for r in rows:
        w.writerow([kr(r[k]) if k in ("price", "best_price", "diff") else
                    (f"{r[k]:.1f}".replace(".", ",") if k == "diff_pct" and r[k] is not None else r[k] or "")
                    for k, _ in COLUMNS])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def write_xlsx(rows):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None
    wb = Workbook()
    ws = wb.active
    ws.title = "Prissammenligning"
    ws.append([h for _, h in COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F4F3E")
    red, green = PatternFill("solid", fgColor="F8D7DA"), PatternFill("solid", fgColor="D4EDDA")
    for r in rows:
        ws.append([(round(r[k], 1) if k == "diff_pct" and r[k] is not None else r[k]) for k, _ in COLUMNS])
        fill = red if r["status"] == "Vi er dyrere" else green if r["status"] == "Vi er billigst" else None
        if fill:
            ws.cell(ws.max_row, 11).fill = fill
    for col in (5, 6, 8):
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            for c in cell:
                c.number_format = '#,##0.00 "kr."'
    widths = [18, 45, 14, 15, 13, 13, 22, 12, 10, 8, 18, 40, 40]
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def write_html(rows, date):
    e = html.escape
    by_sup = {}
    for r in rows:
        by_sup.setdefault(r["supplier"], []).append(r)
    total = len(rows)
    dyrere = sum(r["status"] == "Vi er dyrere" for r in rows)
    billigst = sum(r["status"] in ("Vi er billigst", "Samme pris") for r in rows)
    ikke = total - dyrere - billigst
    out = [f"""<html><body style="font-family:Arial,sans-serif;font-size:13px;color:#222">
<h2 style="color:#2F4F3E">Prissammenligning – {e(SHOP_HOST)} – {date:%d-%m-%Y}</h2>
<p><b>{total}</b> produkter tjekket · <b style="color:#b02a37">{dyrere}</b> dyrere end laveste pris online ·
<b style="color:#1e7e34">{billigst}</b> billigst/samme pris · {ikke} ikke fundet eller uden pris.</p>
<p>Sorteret efter leverandør ({e(', '.join(FIRST_SUPPLIERS))} først). Inden for hver leverandør står de varer
hvor vi er dyrest øverst. Hele listen findes også som Excel-fil.</p>
<h3>Oversigt pr. leverandør</h3>
<table cellpadding="4" cellspacing="0" border="1" style="border-collapse:collapse;border-color:#ccc">
<tr style="background:#2F4F3E;color:#fff"><th>Leverandør</th><th>Produkter</th><th>Vi er dyrere</th>
<th>Vi er billigst</th><th>Ikke fundet</th></tr>"""]
    for sup, rs in by_sup.items():
        d = sum(r["status"] == "Vi er dyrere" for r in rs)
        b = sum(r["status"] in ("Vi er billigst", "Samme pris") for r in rs)
        out.append(f"<tr><td>{e(sup)}</td><td align=right>{len(rs)}</td><td align=right>{d}</td>"
                   f"<td align=right>{b}</td><td align=right>{len(rs) - d - b}</td></tr>")
    out.append("</table>")
    for sup, rs in by_sup.items():
        out.append(f'<h3 style="color:#2F4F3E;margin-top:24px">{e(sup)}</h3>'
                   '<table cellpadding="4" cellspacing="0" border="1" style="border-collapse:collapse;border-color:#ccc">'
                   '<tr style="background:#eee"><th align=left>Produkt</th><th>Vores pris</th><th>Laveste online</th>'
                   '<th align=left>Butik</th><th>Forskel</th><th align=left>Status</th></tr>')
        for r in rs:
            color = "#b02a37" if r["status"] == "Vi er dyrere" else "#1e7e34" if r["status"] == "Vi er billigst" else "#666"
            pct = f" ({r['diff_pct']:+.1f}%)".replace(".", ",") if r["diff_pct"] is not None else ""
            shop = e(r["best_shop"])
            if r["best_url"]:
                shop = f'<a href="{e(r["best_url"])}">{shop}</a>'
            out.append(f'<tr><td><a href="{e(r["url"] or "")}">{e(r["name"])}</a></td>'
                       f'<td align=right>{kr(r["price"])}</td><td align=right>{kr(r["best_price"])}</td>'
                       f'<td>{shop}</td><td align=right>{kr(r["diff"])}{pct}</td>'
                       f'<td style="color:{color}">{e(r["status"])}</td></tr>')
        out.append("</table>")
    out.append('<p style="color:#888;margin-top:24px">Priserne er fundet automatisk og kan indeholde forkerte '
               'match (fx en anden størrelse eller farve). Tjek linket før prisen ændres.</p></body></html>')
    return "\n".join(out)


def md_cell(v):
    return str(v or "").replace("|", "\\|").replace("\n", " ")


def write_markdown(rows, date):
    by_sup = {}
    for r in rows:
        by_sup.setdefault(r["supplier"], []).append(r)
    dyrere = sum(r["status"] == "Vi er dyrere" for r in rows)
    billigst = sum(r["status"] in ("Vi er billigst", "Samme pris") for r in rows)
    out = [f"# Prissammenligning – {SHOP_HOST} – {date:%d-%m-%Y}", "",
           f"**{len(rows)}** produkter tjekket · 🔴 **{dyrere}** dyrere end laveste pris online · "
           f"🟢 **{billigst}** billigst/samme pris · {len(rows) - dyrere - billigst} ikke fundet eller uden pris.", "",
           f"Sorteret efter leverandør ({', '.join(FIRST_SUPPLIERS)} først). Inden for hver leverandør står de varer, "
           "hvor vi er dyrest, øverst. Hele listen findes også som [Excel](seneste.xlsx) og [CSV](seneste.csv).", "",
           "## Oversigt pr. leverandør", "",
           "| Leverandør | Produkter | Vi er dyrere | Vi er billigst | Ikke fundet |", "|---|--:|--:|--:|--:|"]
    for sup, rs in by_sup.items():
        d = sum(r["status"] == "Vi er dyrere" for r in rs)
        b = sum(r["status"] in ("Vi er billigst", "Samme pris") for r in rs)
        anchor = re.sub(r"[^\w\- ]", "", sup.lower()).replace(" ", "-")
        out.append(f"| [{md_cell(sup)}](#{anchor}) | {len(rs)} | {d} | {b} | {len(rs) - d - b} |")
    icons = {"Vi er dyrere": "🔴", "Vi er billigst": "🟢", "Samme pris": "🟢"}
    for sup, rs in by_sup.items():
        out += ["", f"## {md_cell(sup)}", "",
                "| Produkt | Vores pris | Laveste online | Butik | Forskel | Status |", "|---|--:|--:|---|--:|---|"]
        for r in rs:
            name = f"[{md_cell(r['name'])}]({r['url']})" if r["url"] else md_cell(r["name"])
            shop = f"[{md_cell(r['best_shop'])}]({r['best_url']})" if r["best_url"] else md_cell(r["best_shop"])
            pct = f" ({r['diff_pct']:+.1f}%)".replace(".", ",") if r["diff_pct"] is not None else ""
            out.append(f"| {name} | {kr(r['price'])} | {kr(r['best_price'])} | {shop} | {kr(r['diff'])}{pct} | "
                       f"{icons.get(r['status'], '⚪')} {r['status']} |")
    out += ["", "_Priserne er fundet automatisk og kan indeholde forkerte match (fx en anden størrelse eller farve). "
            "Tjek linket, før prisen ændres._", ""]
    return "\n".join(out)


def send_mail(subject, html_body, attachments):
    host = os.environ.get("SMTP_HOST")
    if not host:
        raise SystemExit("SMTP_HOST mangler – sæt DRY_RUN=1 for kun at gemme rapporten.")
    port = int(os.environ.get("SMTP_PORT", "587"))
    user, pwd = os.environ.get("SMTP_USER"), os.environ.get("SMTP_PASSWORD")
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = os.environ.get("MAIL_FROM") or user or MAIL_TO
    msg["To"] = MAIL_TO
    msg.set_content("Rapporten vises bedst i HTML. Se også vedhæftet Excel-fil.")
    msg.add_alternative(html_body, subtype="html")
    for name, data, maintype, subtype in attachments:
        msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=name)
    ctx = ssl.create_default_context()
    if port == 465:
        with smtplib.SMTP_SSL(host, port, context=ctx, timeout=60) as s:
            if user:
                s.login(user, pwd)
            s.send_message(msg)
    else:
        with smtplib.SMTP(host, port, timeout=60) as s:
            s.starttls(context=ctx)
            if user:
                s.login(user, pwd)
            s.send_message(msg)
    print(f"Mail sendt til {MAIL_TO}", file=sys.stderr)


def main():
    date = dt.date.today()
    if os.environ.get("PRODUCTS_FILE"):  # til test uden adgang til shoppen
        with open(os.environ["PRODUCTS_FILE"], encoding="utf-8") as f:
            products = [product(**p) for p in json.load(f)]
    else:
        products = fetch_products()
    for p in products:
        p["supplier"] = guess_supplier(p)
    products.sort(key=lambda p: supplier_key(p["supplier"]))
    if MAX_PRODUCTS:
        products = products[:MAX_PRODUCTS]
    rows = build_rows(products)

    stamp = f"{date:%Y-%m-%d}"
    html_body = write_html(rows, date)
    files = {"md": write_markdown(rows, date).encode("utf-8"), "html": html_body.encode("utf-8"),
             "csv": write_csv(rows)}
    xlsx = write_xlsx(rows)
    if xlsx:
        files["xlsx"] = xlsx

    os.makedirs(os.path.join(OUT_DIR, "arkiv"), exist_ok=True)
    for ext, data in files.items():
        for path in (os.path.join(OUT_DIR, f"seneste.{ext}"), os.path.join(OUT_DIR, "arkiv", f"{stamp}.{ext}")):
            with open(path, "wb") as f:
                f.write(data)
    print(f"Rapport gemt i {OUT_DIR}/seneste.* og arkiv/{stamp}.*", file=sys.stderr)

    if os.environ.get("SMTP_HOST") and os.environ.get("DRY_RUN") != "1":
        dyrere = sum(r["status"] == "Vi er dyrere" for r in rows)
        subject = f"Prissammenligning {date:%d-%m-%Y}: {dyrere} af {len(rows)} varer er dyrere end online"
        attachments = [(f"prissammenligning-{stamp}.csv", files["csv"], "text", "csv")]
        if xlsx:
            attachments.insert(0, (f"prissammenligning-{stamp}.xlsx", xlsx, "application",
                                   "vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
        send_mail(subject, html_body, attachments)


if __name__ == "__main__":
    main()
